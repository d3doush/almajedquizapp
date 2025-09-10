from flask import Blueprint, jsonify, request, session, send_file
from src.models.employee import Score, Employee, db
from datetime import datetime, date
from sqlalchemy import func, and_
import openpyxl
from io import BytesIO
import os

results_bp = Blueprint('results', __name__)

@results_bp.route('/scores', methods=['GET'])
def get_scores():
    if 'user_type' not in session:
        return jsonify({'success': False, 'message': 'غير مصرح لك بالوصول'}), 401
    
    # Get filter parameters
    date_filter = request.args.get('date_filter', 'month')  # month, year
    scope_filter = request.args.get('scope_filter', 'employee')  # employee, store, city, region, all
    
    user_type = session['user_type']
    
    # Base query
    query = Score.query
    
    # Date filtering
    today = date.today()
    if date_filter == 'month':
        query = query.filter(
            func.extract('year', Score.date) == today.year,
            func.extract('month', Score.date) == today.month
        )
    elif date_filter == 'year':
        query = query.filter(func.extract('year', Score.date) == today.year)
    
    # Scope filtering for employees
    if user_type == 'employee':
        employee_name = session['employee_name']
        
        if scope_filter == 'employee':
            query = query.filter(Score.employee_name == employee_name)
        elif scope_filter == 'store':
            store_name = session['employee_store']
            # Get all employees from same store
            store_employees = Employee.query.filter_by(store=store_name).all()
            employee_names = [emp.name for emp in store_employees]
            query = query.filter(Score.employee_name.in_(employee_names))
        elif scope_filter == 'city':
            city_name = session['employee_city']
            city_employees = Employee.query.filter_by(city=city_name).all()
            employee_names = [emp.name for emp in city_employees]
            query = query.filter(Score.employee_name.in_(employee_names))
        elif scope_filter == 'region':
            region_name = session['employee_region']
            region_employees = Employee.query.filter_by(region=region_name).all()
            employee_names = [emp.name for emp in region_employees]
            query = query.filter(Score.employee_name.in_(employee_names))
        # 'all' means no additional filtering
    
    # For admin, apply scope filtering differently
    elif user_type == 'admin':
        if scope_filter == 'store':
            store_name = request.args.get('store_name')
            if store_name:
                store_employees = Employee.query.filter_by(store=store_name).all()
                employee_names = [emp.name for emp in store_employees]
                query = query.filter(Score.employee_name.in_(employee_names))
        elif scope_filter == 'city':
            city_name = request.args.get('city_name')
            if city_name:
                city_employees = Employee.query.filter_by(city=city_name).all()
                employee_names = [emp.name for emp in city_employees]
                query = query.filter(Score.employee_name.in_(employee_names))
        elif scope_filter == 'region':
            region_name = request.args.get('region_name')
            if region_name:
                region_employees = Employee.query.filter_by(region=region_name).all()
                employee_names = [emp.name for emp in region_employees]
                query = query.filter(Score.employee_name.in_(employee_names))
    
    # Execute query and group by employee
    scores = query.all()
    
    # Group scores by employee and calculate totals
    employee_scores = {}
    for score in scores:
        if score.employee_name not in employee_scores:
            employee_scores[score.employee_name] = {
                'employee_name': score.employee_name,
                'store_name': score.store_name,
                'total_score': 0,
                'test_count': 0,
                'average_score': 0
            }
        employee_scores[score.employee_name]['total_score'] += score.score
        employee_scores[score.employee_name]['test_count'] += 1
    
    # Calculate averages and sort
    result_list = []
    for emp_data in employee_scores.values():
        if emp_data['test_count'] > 0:
            emp_data['average_score'] = round(emp_data['total_score'] / emp_data['test_count'], 2)
        result_list.append(emp_data)
    
    # Sort by average score (descending)
    result_list.sort(key=lambda x: x['average_score'], reverse=True)
    
    return jsonify({
        'success': True,
        'scores': result_list,
        'filters': {
            'date_filter': date_filter,
            'scope_filter': scope_filter
        }
    })

@results_bp.route('/export-excel', methods=['POST'])
def export_excel():
    if 'user_type' not in session or session['user_type'] != 'admin':
        return jsonify({'success': False, 'message': 'غير مصرح لك بالوصول'}), 401
    
    data = request.json
    scores_data = data.get('scores', [])
    
    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'النتائج'
    
    # Headers
    headers = ['اسم الموظف', 'اسم المعرض', 'إجمالي النقاط', 'عدد الاختبارات', 'المتوسط']
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    
    # Data
    for row, score in enumerate(scores_data, 2):
        ws.cell(row=row, column=1, value=score['employee_name'])
        ws.cell(row=row, column=2, value=score['store_name'])
        ws.cell(row=row, column=3, value=score['total_score'])
        ws.cell(row=row, column=4, value=score['test_count'])
        ws.cell(row=row, column=5, value=score['average_score'])
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        as_attachment=True,
        download_name=f'results_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@results_bp.route('/locations', methods=['GET'])
def get_locations():
    """Get available regions, cities, and stores for filtering"""
    if 'user_type' not in session:
        return jsonify({'success': False, 'message': 'غير مصرح لك بالوصول'}), 401
    
    employees = Employee.query.all()
    
    regions = list(set([emp.region for emp in employees]))
    cities = list(set([emp.city for emp in employees]))
    stores = list(set([emp.store for emp in employees]))
    
    return jsonify({
        'success': True,
        'regions': regions,
        'cities': cities,
        'stores': stores
    })

