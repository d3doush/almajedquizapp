
import openpyxl

def create_excel_template(file_name='training_data.xlsx'):
    workbook = openpyxl.Workbook()

    # Sheet: staff
    staff_sheet = workbook.active
    staff_sheet.title = 'staff'
    staff_sheet['A1'] = 'المنطقة'
    staff_sheet['B1'] = 'المدينة'
    staff_sheet['C1'] = 'المعرض'
    staff_sheet['D1'] = 'اسم الموظف'
    staff_sheet['E1'] = 'كلمة المرور'

    # Sheet: تدريباتي
    training_sheet = workbook.create_sheet('تدريباتي')
    training_sheet['A1'] = 'اسم العطر'
    training_sheet['B1'] = 'السؤال'
    training_sheet['C1'] = 'إجابة 1'
    training_sheet['D1'] = 'إجابة 2'
    training_sheet['E1'] = 'إجابة 3'
    training_sheet['F1'] = 'إجابة 4'
    training_sheet['G1'] = 'الإجابة الصحيحة'

    # Sheet: علامتي
    scores_sheet = workbook.create_sheet('علامتي')
    scores_sheet['A1'] = 'تاريخ الإجابة'
    scores_sheet['B1'] = 'اسم الموظف'
    scores_sheet['C1'] = 'اسم المعرض'
    scores_sheet['D1'] = 'النتيجة'

    workbook.save(file_name)
    print(f'Excel file \'{file_name}\' created successfully.')

if __name__ == '__main__':
    create_excel_template()


